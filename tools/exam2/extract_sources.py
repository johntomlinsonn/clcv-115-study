# Extract text (incl. speaker notes) from every .pptx/.docx under a folder.
# Usage: python3 -I tools/exam2/extract_sources.py <extracted-zip-dir> <out-dir>
import sys, os, glob
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import docx
raw, out = sys.argv[1], sys.argv[2]
def shapes_text(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes_text(sh.shapes); continue
        if sh.has_text_frame:
            t = "\n".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
            if t: yield t
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows: yield " | ".join(c.text for c in r.cells)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: yield "[IMAGE]"
for f in sorted(glob.glob(raw + "/**/*.pptx", recursive=True)):
    p = Presentation(f); lines = []
    for i, s in enumerate(p.slides, 1):
        lines.append(f"\n--- Slide {i} ---")
        lines += list(shapes_text(s.shapes))
        if s.has_notes_slide:
            n = s.notes_slide.notes_text_frame.text.strip()
            if n: lines.append("NOTES: " + n)
    open(os.path.join(out, os.path.basename(f) + ".txt"), "w").write("\n".join(lines))
for f in sorted(glob.glob(raw + "/**/*.docx", recursive=True)):
    d = docx.Document(f); lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows: lines.append(" | ".join(c.text.strip() for c in r.cells))
    open(os.path.join(out, os.path.basename(f) + ".txt"), "w").write("\n".join(lines))
