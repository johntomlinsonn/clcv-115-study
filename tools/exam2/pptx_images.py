# Save the pictures on chosen slides of a .pptx as s<slide>_<n>.<ext>.
# Usage: python3 -I tools/exam2/pptx_images.py <file.pptx> <out-dir> 21,60,92
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
p = Presentation(sys.argv[1]); out = sys.argv[2]; want = set(int(x) for x in sys.argv[3].split(","))
def pics(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP: yield from pics(sh.shapes)
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE: yield sh
for i, s in enumerate(p.slides, 1):
    if i not in want: continue
    for n, sh in enumerate(pics(s.shapes), 1):
        open(f"{out}/s{i:02d}_{n}.{sh.image.ext}", "wb").write(sh.image.blob)
